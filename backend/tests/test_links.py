"""Links: kept from the original file, and clickable wherever the resume appears."""

import io

import docx
import pymupdf
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml.shared import OxmlElement, qn

from app.rendering.render import page_images, render_html, render_pdf
from app.schemas.resume import ResumeData
from app.services.extract import DOCX, LINKS_HEADING, PDF, extract_text
from app.services.parse_resume import PCertification, PLink, PProject, to_resume_data
from tests.test_parse_resume import SOURCE, parsed

CREDLY = "https://www.credly.com/badges/3f1c2a9e-77b0-4c1e-9d55-0b2f6a1e8c41/public_url"

WITH_CERT = ResumeData.model_validate(
    {
        "basics": {"name": "Asha Rao", "email": "asha@example.com", "phone": "+91 90000 00000"},
        "certifications": [
            {"name": "AWS Certified Developer", "issuer": "Amazon", "date": "2023", "url": CREDLY}
        ],
    }
)


# --- reading links out of the original file ---------------------------------------------


def linked_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Asha Rao — Backend Engineer, asha@example.com")
    page.insert_text((72, 100), "Certifications")
    page.insert_text((72, 120), "AWS Certified Developer (2023)")
    page.insert_text((72, 140), "LinkedIn")
    rect = pymupdf.Rect(70, 108, 260, 124)
    page.insert_link({"kind": pymupdf.LINK_URI, "from": rect, "uri": CREDLY})
    page.insert_link(
        {
            "kind": pymupdf.LINK_URI,
            "from": pymupdf.Rect(70, 128, 130, 144),
            "uri": "https://linkedin.com/in/asha",
        }
    )
    return doc.tobytes()


def test_links_hidden_behind_text_in_a_pdf_are_listed():
    text = extract_text(linked_pdf(), PDF)
    listed = text.split(LINKS_HEADING, 1)[1]
    assert f"AWS Certified Developer (2023) → {CREDLY}" in listed
    assert "LinkedIn → https://linkedin.com/in/asha" in listed


def test_links_hidden_behind_text_in_a_word_file_are_listed():
    document = docx.Document()
    document.add_paragraph("Asha Rao, Backend Engineer, asha@example.com, Pune, India")
    paragraph = document.add_paragraph("Certifications: ")
    rid = paragraph.part.relate_to(CREDLY, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "AWS Certified Developer"
    run.append(t)
    link.append(run)
    paragraph._p.append(link)
    buf = io.BytesIO()
    document.save(buf)
    text = extract_text(buf.getvalue(), DOCX)
    assert f"AWS Certified Developer → {CREDLY}" in text


def test_a_file_without_hidden_links_has_no_links_section():
    from tests import fixtures

    assert LINKS_HEADING not in extract_text(fixtures.resume_docx(), DOCX)


# --- the parse keeps only addresses the file has ---------------------------------------------


def test_a_certificates_link_from_the_file_is_kept():
    source = SOURCE + f"\n{LINKS_HEADING}\n- AWS Certified Developer → {CREDLY}\n"
    data, _ = to_resume_data(
        parsed(
            certifications=[
                PCertification(name="AWS Certified Developer", issuer="", date="2023", url=CREDLY)
            ]
        ),
        source,
    )
    assert data.certifications[0].url == CREDLY


def test_an_address_the_file_doesnt_have_is_dropped():
    data, _ = to_resume_data(
        parsed(
            certifications=[
                PCertification(name="AWS", issuer="", date="", url="https://credly.com/made-up")
            ],
            projects=[
                PProject(
                    name="ledgerkit", url="github.com/asha/ledgerkit", start="", end="", bullets=[]
                )
            ],
        ),
        SOURCE,
    )
    assert data.certifications[0].url == "" and data.projects[0].url == ""


def test_a_written_out_address_matches_however_it_is_spelled():
    source = SOURCE + "\nGitHub: github.com/asha-rao\n"
    basics = parsed().basics.model_copy(
        update={"links": [PLink(label="GitHub", url="https://www.github.com/asha-rao/")]}
    )
    data, _ = to_resume_data(parsed(basics=basics), source)
    assert data.basics.links[0].url == "https://www.github.com/asha-rao/"


# --- clickable in the resume -------------------------------------------------------------------


def test_a_certificate_links_to_its_credential():
    html = render_html(WITH_CERT, "classic")
    assert f'<a href="{CREDLY}"><strong>AWS Certified Developer</strong></a>' in html
    assert ">credly.com/…</a>" in html  # short to read; the full address in the link


def test_every_link_is_clickable_in_the_pdf():
    pdf = render_pdf(render_html(WITH_CERT, "classic")).content
    with pymupdf.open(stream=pdf, filetype="pdf") as doc:
        uris = {link["uri"] for link in doc[0].get_links()}
    assert {CREDLY, "mailto:asha@example.com", "tel:+919000000000"} <= uris


def test_the_preview_knows_where_the_links_are():
    [page] = page_images(render_pdf(render_html(WITH_CERT, "classic")).content)
    credly = [link for link in page.links if link["url"] == CREDLY]
    assert credly and all(0 <= link["x"] < 1 and 0 <= link["y"] < 1 for link in credly)


def test_only_web_mail_and_phone_links_are_offered_in_the_preview():
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_link(
        {"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(0, 0, 50, 50), "uri": "javascript:alert(1)"}
    )
    assert page_images(doc.tobytes())[0].links == []


# --- choosing what's in the header -------------------------------------------------------


def test_header_items_can_be_left_off_a_resume():
    from app.schemas.layout import Layout

    data = WITH_CERT.model_copy(deep=True)
    data.basics.headline = "Backend Engineer"
    data.basics.location = "Pune"
    data.basics.links = [
        {"id": "lnk_li", "label": "LinkedIn", "url": "linkedin.com/in/asha"},
        {"id": "lnk_gh", "label": "GitHub", "url": "github.com/asha"},
    ]
    data = ResumeData.model_validate(data.model_dump())
    html = render_html(data, "classic", Layout(hidden_header=["phone", "headline", "lnk_li"]))
    assert "+91 90000 00000" not in html and "Backend Engineer" not in html
    assert "linkedin.com/in/asha" not in html
    assert "github.com/asha" in html and "asha@example.com" in html and "Pune" in html


def test_the_cover_letter_leaves_off_the_same_header_items():
    from app.rendering.render import render_letter_html
    from app.schemas.layout import Layout

    html = render_letter_html(
        WITH_CERT, "Hello.", "Acme", "classic", Layout(hidden_header=["phone"])
    )
    assert "+91 90000 00000" not in html and "asha@example.com" in html
