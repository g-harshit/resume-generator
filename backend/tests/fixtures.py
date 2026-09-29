"""Resume files built in code, so the tests need no binary fixtures and every person
in them is invented."""

import io

import docx
import pymupdf

MAIN = [
    "Experience",
    "Backend Engineer, Finlytics",
    "Designed a double-entry ledger service in Go handling 2.1M transactions a day.",
    "Split the settlement monolith into gRPC services on AWS.",
    "Software Engineer, Cartwheel",
    "Built the order-routing service in Go serving 40 warehouses.",
]
SIDEBAR = ["Asha Rao", "asha.rao@example.com", "Skills", "Go, PostgreSQL, Kafka, gRPC"]


def two_column_pdf() -> bytes:
    """A sidebar on the left, the main column on the right, under a full-width title."""
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_textbox(pymupdf.Rect(40, 30, 555, 60), "Asha Rao - Backend Engineer", fontsize=14)
    y = 90
    for line in SIDEBAR:
        page.insert_textbox(pymupdf.Rect(40, y, 190, y + 40), line, fontsize=10)
        y += 50
    y = 90
    for line in MAIN:
        page.insert_textbox(pymupdf.Rect(220, y, 555, y + 40), line, fontsize=10)
        y += 50
    return doc.tobytes()


def dated_single_column_pdf() -> bytes:
    """One column with dates right-aligned on the same line as each role."""
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    rows = [
        ("Backend Engineer, Finlytics", "Mar 2021 - Present"),
        ("Designed a double-entry ledger service in Go handling 2.1M transactions.", None),
        ("Software Engineer, Cartwheel", "Jul 2019 - Feb 2021"),
        ("Built the order-routing service in Go serving 40 warehouses.", None),
    ]
    y = 60
    for text, date in rows:
        page.insert_textbox(pymupdf.Rect(40, y, 420, y + 30), text, fontsize=10)
        if date:
            page.insert_textbox(pymupdf.Rect(450, y, 555, y + 30), date, fontsize=10)
        y += 40
    return doc.tobytes()


def blank_pdf() -> bytes:
    doc = pymupdf.open()
    doc.new_page()
    return doc.tobytes()


def resume_docx(*, contact_in_header: bool = True) -> bytes:
    document = docx.Document()
    if contact_in_header:
        document.sections[0].header.paragraphs[0].text = "Asha Rao · asha.rao@example.com"
    document.add_heading("Experience", level=1)
    document.add_paragraph("Backend Engineer, Finlytics — Mar 2021 to Present")
    document.add_paragraph("Designed a double-entry ledger service in Go.", style="List Bullet")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Go, PostgreSQL, Kafka, gRPC"
    document.add_paragraph("Education: B.Tech Computer Science, 2019")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
