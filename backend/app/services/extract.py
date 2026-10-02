"""Uploaded file → plain text, in reading order.

The file type is decided from its bytes, never from the name or the browser's MIME
type. Anything we can't read raises `ExtractionError` with a message for the user.
"""

import io
import re
import zipfile
from dataclasses import dataclass

import docx
import pymupdf

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# A scanned (image-only) PDF yields almost no text; below this we say so rather than
# sending the model an empty page.
MIN_TEXT_CHARS = 80


class ExtractionError(Exception):
    pass


def sniff_type(data: bytes) -> str | None:
    if data.startswith(b"%PDF-"):
        return PDF
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" in z.namelist():
                    return DOCX
        except zipfile.BadZipFile:
            return None
    return None


def extract_text(data: bytes, mime: str) -> str:
    text = _pdf_text(data) if mime == PDF else _docx_text(data)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < MIN_TEXT_CHARS:
        raise ExtractionError(
            "We couldn't find any text in this file. If it's a scan or a photo, upload the "
            "Word version or a PDF exported from your editor instead."
        )
    return text


# --- PDF -----------------------------------------------------------------------


@dataclass
class _Block:
    x0: float
    y0: float
    x1: float
    y1: float
    text: str


# A block this much of the page wide is a header or a full-width section, not part of
# a column. A two-column resume's main column is often 60-65% of the page, so keep this high.
_WIDE = 0.75
# A side column must hold at least this share of a segment's text to count as a real
# column. Below it, it's a strip of right-aligned dates that belongs to its lines.
_MIN_COLUMN_SHARE = 0.15
# Left edges this far apart (share of page width) start different columns.
_COLUMN_GAP = 0.15


def _pdf_text(data: bytes) -> str:
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ExtractionError("This PDF looks damaged. Try exporting it again.") from exc
    if doc.needs_pass:
        raise ExtractionError("This PDF is password-protected. Upload a copy without a password.")
    with doc:
        text = "\n\n".join(_page_text(page) for page in doc)
        links = [
            (page.get_textbox(link["from"] + (-1, -1, 1, 1)), link["uri"])
            for page in doc
            for link in page.get_links()
            if link.get("uri")
        ]
    return text + links_section(links)


# Where a file hides a link behind text (a certificate's name linked to its credential,
# "LinkedIn" linked to a profile), the text alone loses the address. These are listed
# after the text so the parser can attach each to its item.
LINKS_HEADING = "Links in the file (the linked text → where it points):"


def links_section(links: list[tuple[str, str]]) -> str:
    seen: set[str] = set()
    lines = []
    for text, url in links:
        url = url.strip()
        if not url or url in seen or url.lower().startswith(("mailto:", "tel:")):
            continue  # an email or phone is already in the text as itself
        seen.add(url)
        anchor = " ".join(text.split())[:120]
        lines.append(f"- {anchor or '(no text)'} → {url}")
    return f"\n\n{LINKS_HEADING}\n" + "\n".join(lines) if lines else ""


def _page_text(page) -> str:
    width = page.rect.width
    blocks = sorted(
        (
            _Block(b[0], b[1], b[2], b[3], b[4].strip())
            for b in page.get_text("blocks")
            if b[6] == 0 and b[4].strip()
        ),
        key=lambda b: (b.y0, b.x0),
    )

    # Full-width blocks split the page into segments; columns are found per segment,
    # so a sidebar that starts below a full-width header is still read as a column.
    out: list[str] = []
    run: list[_Block] = []
    for block in blocks:
        if (block.x1 - block.x0) >= _WIDE * width:
            out += _segment_text(run, width)
            run = []
            out.append(block.text)
        else:
            run.append(block)
    out += _segment_text(run, width)
    return "\n".join(out)


def _segment_text(blocks: list[_Block], width: float) -> list[str]:
    if not blocks:
        return []
    columns = _columns(blocks, width)
    total = sum(len(b.text) for b in blocks)
    real_columns = len(columns) > 1 and all(
        sum(len(b.text) for b in col) >= _MIN_COLUMN_SHARE * total for col in columns
    )
    if not real_columns:
        # One column (perhaps with dates on the right): read line by line.
        return [b.text for b in sorted(blocks, key=lambda b: (round(b.y0), b.x0))]
    return [b.text for col in columns for b in sorted(col, key=lambda b: b.y0)]


def _columns(blocks: list[_Block], width: float) -> list[list[_Block]]:
    """Group blocks by where they start, left to right.

    Left edges, not overlap: a name or heading at the top can run across the gutter,
    and overlap-grouping would then merge both columns into one. Column starts are far
    apart; an indented bullet starts only a little right of its column."""
    groups: list[list[_Block]] = []
    for block in sorted(blocks, key=lambda b: b.x0):
        if groups and block.x0 - groups[-1][-1].x0 < _COLUMN_GAP * width:
            groups[-1].append(block)
        else:
            groups.append([block])
    return groups


# --- DOCX ----------------------------------------------------------------------


def _docx_text(data: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("This Word file looks damaged. Try saving it again.") from exc

    lines: list[str] = []
    # Contact details often sit in the page header, where a plain paragraph walk
    # would never find them.
    for section in document.sections:
        lines += [p.text for p in section.header.paragraphs]

    # Paragraphs and table cells in document order (resumes often lay out with tables).
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append(_xml_text(child))
        elif tag == "tbl":
            lines += [_xml_text(cell) for cell in child.iter() if cell.tag.endswith("}tc")]

    for section in document.sections:
        lines += [p.text for p in section.footer.paragraphs]

    links = _docx_links(document.part)
    for section in document.sections:
        links += _docx_links(section.header.part) + _docx_links(section.footer.part)
    return "\n".join(line for line in lines if line.strip()) + links_section(links)


def _docx_links(part) -> list[tuple[str, str]]:
    """(text, address) for each hyperlink in one part of a Word file (body, header…)."""
    out = []
    for el in part.element.iter():
        if not el.tag.endswith("}hyperlink"):
            continue
        rid = el.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        rel = part.rels.get(rid) if rid else None
        if rel is not None and rel.is_external:
            out.append((_xml_text(el), rel.target_ref))
    return out


def _xml_text(element) -> str:
    return "".join(t.text or "" for t in element.iter() if t.tag.endswith("}t"))
